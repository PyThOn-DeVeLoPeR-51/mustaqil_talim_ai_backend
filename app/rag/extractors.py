from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

import fitz
from docx import Document


class DocumentExtractionError(ValueError):
    """Fayldan foydali matn ajratib bo‘lmaganda ko‘tariladi."""


@dataclass(frozen=True)
class ExtractedBlock:
    text: str
    page_number: int | None = None
    section_title: str | None = None


@dataclass(frozen=True)
class ExtractedDocument:
    blocks: list[ExtractedBlock]
    page_count: int | None
    metadata: dict[str, object]


def _normalize_text(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]

    paragraphs: list[str] = []
    current: list[str] = []
    for line in lines:
        if line:
            current.append(line)
            continue
        if current:
            paragraphs.append(" ".join(current))
            current = []
    if current:
        paragraphs.append(" ".join(current))

    return "\n\n".join(paragraphs).strip()


def extract_pdf(path: str | Path) -> ExtractedDocument:
    source = Path(path)
    blocks: list[ExtractedBlock] = []

    try:
        with fitz.open(str(source)) as document:
            page_count = document.page_count
            for page_index in range(page_count):
                page = document.load_page(page_index)
                text = _normalize_text(page.get_text("text"))
                if text:
                    blocks.append(
                        ExtractedBlock(
                            text=text,
                            page_number=page_index + 1,
                        )
                    )
    except Exception as exc:  # pragma: no cover - kutubxona xatolari turlicha
        raise DocumentExtractionError(
            "PDF faylni o‘qib bo‘lmadi yoki fayl buzilgan."
        ) from exc

    if not blocks:
        raise DocumentExtractionError(
            "PDF ichidan matn topilmadi. Fayl skanerlangan rasm bo‘lishi mumkin; "
            "OCR keyingi bosqichda alohida qo‘shiladi."
        )

    return ExtractedDocument(
        blocks=blocks,
        page_count=page_count,
        metadata={
            "extractor": "pymupdf",
            "pages_with_text": len(blocks),
        },
    )


def extract_docx(path: str | Path) -> ExtractedDocument:
    source = Path(path)
    blocks: list[ExtractedBlock] = []
    current_section: str | None = None

    try:
        document = Document(str(source))
    except Exception as exc:  # pragma: no cover
        raise DocumentExtractionError(
            "DOCX faylni o‘qib bo‘lmadi yoki fayl buzilgan."
        ) from exc

    for paragraph in document.paragraphs:
        text = _normalize_text(paragraph.text)
        if not text:
            continue

        style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style_name.startswith("heading") or style_name.startswith("title"):
            current_section = text[:500]
            # Sarlavhani ham qidiruv kontekstida saqlaymiz.
            blocks.append(
                ExtractedBlock(
                    text=text,
                    section_title=current_section,
                )
            )
        else:
            blocks.append(
                ExtractedBlock(
                    text=text,
                    section_title=current_section,
                )
            )

    # Jadvallar DOCX materiallarda muhim bo‘lishi mumkin. Ularni matn ko‘rinishida
    # hujjat oxiriga qo‘shamiz; python-docx pagination bermaydi.
    for table_index, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [_normalize_text(cell.text) for cell in row.cells]
            cells = [cell for cell in cells if cell]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            blocks.append(
                ExtractedBlock(
                    text="\n".join(rows),
                    section_title=f"Jadval {table_index}",
                )
            )

    if not blocks:
        raise DocumentExtractionError("DOCX fayl ichidan matn topilmadi.")

    return ExtractedDocument(
        blocks=blocks,
        page_count=None,
        metadata={
            "extractor": "python-docx",
            "paragraph_count": len(document.paragraphs),
            "table_count": len(document.tables),
        },
    )


def extract_document(path: str | Path, file_type: str) -> ExtractedDocument:
    normalized_type = file_type.lower().strip()
    if normalized_type == "pdf":
        return extract_pdf(path)
    if normalized_type == "docx":
        return extract_docx(path)
    raise DocumentExtractionError(f"Qo‘llab-quvvatlanmaydigan fayl turi: {file_type}")
