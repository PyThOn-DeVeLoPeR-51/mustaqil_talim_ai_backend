from __future__ import annotations

from datetime import timedelta
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.v1.endpoints.rag import router
from app.core.config import settings
from app.db.base import Base
from app.db.database import get_db
from app.models.rag import RAGDocument, RAGProcessingJob
from app.models.teacher import Teacher
from app.rag.embeddings import EmbeddingProviderError
from app.services.auth_service import get_current_teacher
from app.services.rag_job_service import (
    claim_next_rag_job,
    enqueue_rag_job,
    execute_claimed_rag_job,
    get_teacher_monitoring_summary_values,
    recover_stale_jobs,
    retry_teacher_job,
    utcnow,
)
from app.services.rag_service import create_rag_document_record


def make_docx_bytes(paragraphs: int = 8) -> bytes:
    buffer = BytesIO()
    document = Document()
    document.add_heading("Mustaqil ta’lim", level=1)
    for index in range(paragraphs):
        document.add_paragraph(
            f"{index + 1}-paragraf. Talabaning mustaqil faoliyati reja, nazorat va refleksiya asosida tashkil etiladi. "
            * 3
        )
    document.save(buffer)
    return buffer.getvalue()


class FailingEmbeddingProvider:
    provider_name = "failing"
    model_name = "failing-model"
    dimensions = 4

    def embed_documents(self, texts):
        raise EmbeddingProviderError("test embedding failure")

    def embed_query(self, text):
        raise EmbeddingProviderError("test query failure")


class RAGJobTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_storage_dir = settings.RAG_STORAGE_DIR
        self.original_min_free = settings.RAG_STORAGE_MIN_FREE_MB
        self.original_storage_limit = settings.RAG_TEACHER_STORAGE_LIMIT_MB
        self.original_document_limit = settings.RAG_TEACHER_DOCUMENT_LIMIT
        self.original_retry_base = settings.RAG_JOB_RETRY_BASE_SECONDS
        settings.RAG_STORAGE_DIR = self.temp_dir.name
        settings.RAG_STORAGE_MIN_FREE_MB = 0
        settings.RAG_TEACHER_STORAGE_LIMIT_MB = 100
        settings.RAG_TEACHER_DOCUMENT_LIMIT = 20
        settings.RAG_JOB_RETRY_BASE_SECONDS = 0

        self.engine = create_engine(
            "sqlite+pysqlite:///:memory:",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.factory = sessionmaker(bind=self.engine, expire_on_commit=False)
        self.db: Session = self.factory()
        self.teacher = Teacher(
            first_name="Queue",
            last_name="Teacher",
            email="queue-teacher@example.com",
            password_hash="hash",
        )
        self.db.add(self.teacher)
        self.db.commit()
        self.db.refresh(self.teacher)

    def tearDown(self) -> None:
        self.db.close()
        self.engine.dispose()
        settings.RAG_STORAGE_DIR = self.original_storage_dir
        settings.RAG_STORAGE_MIN_FREE_MB = self.original_min_free
        settings.RAG_TEACHER_STORAGE_LIMIT_MB = self.original_storage_limit
        settings.RAG_TEACHER_DOCUMENT_LIMIT = self.original_document_limit
        settings.RAG_JOB_RETRY_BASE_SECONDS = self.original_retry_base
        self.temp_dir.cleanup()

    def _upload_file(self, data: bytes | None = None):
        from fastapi import UploadFile

        payload = data or make_docx_bytes()
        return UploadFile(
            filename="lesson.docx",
            file=BytesIO(payload)
        )

    def test_ingest_job_is_durable_and_succeeds(self) -> None:
        document = create_rag_document_record(
            self.db,
            self.teacher,
            title="Background lesson",
            task_id=None,
            file=self._upload_file(),
        )
        job = enqueue_rag_job(
            self.db,
            self.teacher,
            document,
            job_type="ingest",
            auto_embed=False,
        )
        self.assertEqual(job.status, "pending")
        claimed = claim_next_rag_job(self.db, "test-worker")
        self.assertIsNotNone(claimed)
        self.assertEqual(claimed.id, job.id)
        completed = execute_claimed_rag_job(self.db, claimed)
        self.assertEqual(completed.status, "succeeded")
        self.assertEqual(completed.progress_percent, 100)
        self.db.refresh(document)
        self.assertEqual(document.status, "ready")
        self.assertGreater(document.chunk_count, 0)
        self.assertGreater(completed.result_json["duration_ms"], 0)

    def test_auto_embed_creates_follow_up_job(self) -> None:
        document = create_rag_document_record(
            self.db,
            self.teacher,
            title="Auto embed lesson",
            task_id=None,
            file=self._upload_file(),
        )
        job = enqueue_rag_job(
            self.db,
            self.teacher,
            document,
            job_type="ingest",
            auto_embed=True,
        )
        claimed = claim_next_rag_job(self.db, "test-worker")
        completed = execute_claimed_rag_job(self.db, claimed)
        self.assertEqual(completed.status, "succeeded")
        next_job_id = completed.result_json["next_job_id"]
        next_job = self.db.query(RAGProcessingJob).filter_by(id=next_job_id).one()
        self.assertEqual(next_job.job_type, "embed")
        self.assertEqual(next_job.status, "pending")

    def test_failed_job_can_be_manually_retried(self) -> None:
        document = RAGDocument(
            teacher_id=self.teacher.id,
            title="Ready document",
            original_filename="ready.docx",
            stored_file_path=str(Path(self.temp_dir.name) / "ready.docx"),
            file_type="docx",
            status="ready",
            chunk_count=1,
        )
        self.db.add(document)
        self.db.commit()
        self.db.refresh(document)
        from app.models.rag import RAGChunk

        self.db.add(RAGChunk(document_id=document.id, chunk_index=0, content="Test chunk"))
        self.db.commit()
        job = enqueue_rag_job(
            self.db,
            self.teacher,
            document,
            job_type="embed",
            max_attempts=1,
        )
        claimed = claim_next_rag_job(self.db, "test-worker")
        failed = execute_claimed_rag_job(
            self.db,
            claimed,
            embedding_provider=FailingEmbeddingProvider(),
        )
        self.assertEqual(failed.status, "failed")
        self.assertIn("test embedding failure", failed.error_message)
        retried = retry_teacher_job(self.db, self.teacher, failed.id)
        self.assertEqual(retried.status, "pending")
        self.assertEqual(retried.attempts, 0)

    def test_stale_running_job_is_recovered(self) -> None:
        document = RAGDocument(
            teacher_id=self.teacher.id,
            title="Stale",
            original_filename="stale.docx",
            stored_file_path="stale.docx",
            file_type="docx",
            status="uploaded",
            chunk_count=0,
        )
        self.db.add(document)
        self.db.flush()
        job = RAGProcessingJob(
            document_id=document.id,
            teacher_id=self.teacher.id,
            job_type="ingest",
            status="running",
            progress_percent=25,
            attempts=1,
            max_attempts=3,
            available_at=utcnow(),
            locked_at=utcnow() - timedelta(minutes=settings.RAG_JOB_STALE_MINUTES + 1),
            locked_by="dead-worker",
        )
        self.db.add(job)
        self.db.commit()
        self.assertEqual(recover_stale_jobs(self.db), 1)
        self.db.refresh(job)
        self.assertEqual(job.status, "pending")
        self.assertIsNone(job.locked_by)

    def test_storage_quota_rejects_and_cleans_file(self) -> None:
        settings.RAG_TEACHER_STORAGE_LIMIT_MB = 0.0001
        with self.assertRaises(Exception) as context:
            create_rag_document_record(
                self.db,
                self.teacher,
                title="Too large for quota",
                task_id=None,
                file=self._upload_file(),
            )
        self.assertEqual(getattr(context.exception, "status_code", None), 507)
        teacher_dir = Path(self.temp_dir.name) / str(self.teacher.id)
        self.assertFalse(teacher_dir.exists() and any(teacher_dir.iterdir()))

    def test_monitoring_summary_counts_jobs_and_documents(self) -> None:
        document = create_rag_document_record(
            self.db,
            self.teacher,
            title="Monitoring lesson",
            task_id=None,
            file=self._upload_file(),
        )
        enqueue_rag_job(self.db, self.teacher, document, job_type="ingest")
        summary = get_teacher_monitoring_summary_values(self.db, self.teacher)
        self.assertEqual(summary["documents_total"], 1)
        self.assertEqual(summary["documents_by_status"]["uploaded"], 1)
        self.assertEqual(summary["jobs_by_status"]["pending"], 1)
        self.assertGreater(summary["storage"]["used_bytes"], 0)


class RAGBackgroundAPITestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_storage_dir = settings.RAG_STORAGE_DIR
        self.original_min_free = settings.RAG_STORAGE_MIN_FREE_MB
        settings.RAG_STORAGE_DIR = self.temp_dir.name
        settings.RAG_STORAGE_MIN_FREE_MB = 0

        self.engine = create_engine(
            "sqlite+pysqlite:///:memory:",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        factory = sessionmaker(bind=self.engine, expire_on_commit=False)
        self.db: Session = factory()
        self.teacher = Teacher(
            first_name="API",
            last_name="Queue",
            email="api-queue@example.com",
            password_hash="hash",
        )
        self.db.add(self.teacher)
        self.db.commit()
        self.db.refresh(self.teacher)

        app = FastAPI()
        app.include_router(router, prefix="/rag")

        def override_get_db():
            yield self.db

        app.dependency_overrides[get_db] = override_get_db
        app.dependency_overrides[get_current_teacher] = lambda: self.teacher
        self.client = TestClient(app)

    def tearDown(self) -> None:
        self.client.close()
        self.db.close()
        self.engine.dispose()
        settings.RAG_STORAGE_DIR = self.original_storage_dir
        settings.RAG_STORAGE_MIN_FREE_MB = self.original_min_free
        self.temp_dir.cleanup()

    def test_background_upload_returns_202_and_job(self) -> None:
        response = self.client.post(
            "/rag/documents/background",
            data={"title": "Async lesson", "auto_embed": "false"},
            files={
                "file": (
                    "lesson.docx",
                    make_docx_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(response.status_code, 202, response.text)
        payload = response.json()
        self.assertEqual(payload["document"]["status"], "uploaded")
        self.assertEqual(payload["job"]["status"], "pending")
        job_response = self.client.get(f"/rag/jobs/{payload['job']['id']}")
        self.assertEqual(job_response.status_code, 200)
        self.assertEqual(job_response.json()["job_type"], "ingest")


if __name__ == "__main__":
    unittest.main()
