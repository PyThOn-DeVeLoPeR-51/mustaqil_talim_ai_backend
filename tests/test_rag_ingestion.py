from __future__ import annotations

from io import BytesIO
from pathlib import Path
import tempfile
import unittest

import fitz
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.v1.endpoints.rag import router
from app.core.config import settings
from app.db.base import Base
from app.db.database import get_db
from app.models.teacher import Teacher
from app.rag.chunker import chunk_blocks
from app.rag.extractors import ExtractedBlock, extract_docx, extract_pdf
from app.services.auth_service import get_current_teacher


def make_pdf_bytes(text: str) -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    data = document.tobytes()
    document.close()
    return data


def make_docx_bytes() -> bytes:
    buffer = BytesIO()
    document = Document()
    document.add_heading("Ortogonal proyeksiya", level=1)
    document.add_paragraph(
        "Ortogonal proyeksiya buyumni o‘zaro perpendikulyar tekisliklarda tasvirlash usulidir."
    )
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Ko‘rinish"
    table.cell(0, 1).text = "Tekislik"
    table.cell(1, 0).text = "Old"
    table.cell(1, 1).text = "Frontal"
    document.save(buffer)
    return buffer.getvalue()


class RAGExtractionTestCase(unittest.TestCase):
    def test_chunker_preserves_page_metadata(self) -> None:
        text = " ".join([f"gap{i}." for i in range(300)])
        chunks = chunk_blocks(
            [ExtractedBlock(text=text, page_number=3, section_title="Mavzu")],
            chunk_size=500,
            overlap=80,
        )
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk.page_number_start == 3 for chunk in chunks))
        self.assertTrue(all(chunk.section_title == "Mavzu" for chunk in chunks))
        self.assertEqual([chunk.chunk_index for chunk in chunks], list(range(len(chunks))))


    def test_docx_style_blocks_are_merged_before_chunking(self) -> None:
        blocks = [
            ExtractedBlock(text="Kredit modul tizimida mustaqil ta’lim.", section_title=None),
            *[
                ExtractedBlock(
                    text=(
                        f"{index}-paragraf. "
                        + "Talabaning mustaqil faoliyati rejalashtirish, nazorat va refleksiya bilan bog‘liq. " * 3
                    ),
                    section_title=None,
                )
                for index in range(1, 13)
            ],
        ]
        chunks = chunk_blocks(blocks, chunk_size=700, overlap=100)

        self.assertLess(len(chunks), len(blocks))
        self.assertIn("Kredit modul tizimida", chunks[0].content)
        self.assertIn("1-paragraf", chunks[0].content)
        self.assertTrue(any(chunk.char_count >= 350 for chunk in chunks[:-1]))

    def test_overlap_starts_on_sentence_boundary(self) -> None:
        text = (
            "Birinchi gap kredit-modul tizimini tushuntiradi. "
            "Ikkinchi gap mustaqil ta’limning ahamiyatini bayon qiladi. "
            "Uchinchi gap talabaga rejalashtirish imkonini beradi. "
            "To‘rtinchi gap baholash va refleksiyani izohlaydi. "
        ) * 8
        chunks = chunk_blocks(
            [ExtractedBlock(text=text, page_number=None, section_title=None)],
            chunk_size=500,
            overlap=120,
        )
        self.assertGreater(len(chunks), 1)
        for chunk in chunks[1:]:
            # Sentence-aware overlap should begin with a complete sentence in
            # this synthetic text, not with a clipped continuation.
            self.assertRegex(
                chunk.content,
                r"^(Birinchi|Ikkinchi|Uchinchi|To‘rtinchi) gap",
            )

    def test_pdf_extraction_returns_page_number(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "lesson.pdf"
            path.write_bytes(make_pdf_bytes("Ortogonal proyeksiya bo‘yicha dars materiali"))
            extracted = extract_pdf(path)
            self.assertEqual(extracted.page_count, 1)
            self.assertEqual(extracted.blocks[0].page_number, 1)
            self.assertIn("Ortogonal proyeksiya", extracted.blocks[0].text)

    def test_docx_extraction_preserves_heading_and_table(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "lesson.docx"
            path.write_bytes(make_docx_bytes())
            extracted = extract_docx(path)
            all_text = "\n".join(block.text for block in extracted.blocks)
            self.assertIn("Ortogonal proyeksiya", all_text)
            self.assertIn("Frontal", all_text)
            self.assertTrue(
                any(block.section_title == "Ortogonal proyeksiya" for block in extracted.blocks)
            )


class RAGAPITestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_storage_dir = settings.RAG_STORAGE_DIR
        self.original_chunk_size = settings.RAG_CHUNK_SIZE_CHARS
        self.original_overlap = settings.RAG_CHUNK_OVERLAP_CHARS
        settings.RAG_STORAGE_DIR = self.temp_dir.name
        settings.RAG_CHUNK_SIZE_CHARS = 500
        settings.RAG_CHUNK_OVERLAP_CHARS = 80

        self.engine = create_engine(
            "sqlite+pysqlite:///:memory:",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        factory = sessionmaker(bind=self.engine, expire_on_commit=False)
        self.db: Session = factory()

        self.teacher = Teacher(
            first_name="RAG",
            last_name="Teacher",
            email="rag-teacher@example.com",
            password_hash="hash",
        )
        self.db.add(self.teacher)
        self.db.commit()
        self.db.refresh(self.teacher)

        app = FastAPI()
        app.include_router(router, prefix="/rag")

        def override_get_db():
            yield self.db

        app.dependency_overrides[get_db] = override_get_db
        app.dependency_overrides[get_current_teacher] = lambda: self.teacher
        self.client = TestClient(app)

    def tearDown(self) -> None:
        self.client.close()
        self.db.close()
        self.engine.dispose()
        settings.RAG_STORAGE_DIR = self.original_storage_dir
        settings.RAG_CHUNK_SIZE_CHARS = self.original_chunk_size
        settings.RAG_CHUNK_OVERLAP_CHARS = self.original_overlap
        self.temp_dir.cleanup()

    def test_pdf_upload_ingestion_list_chunks_and_delete(self) -> None:
        pdf_bytes = make_pdf_bytes(
            "Ortogonal proyeksiya. " + ("Talaba proyeksiyalarni amaliy mashq orqali o‘rganadi. " * 40)
        )
        response = self.client.post(
            "/rag/documents",
            data={"title": "Muhandislik grafikasi ma'ruzasi"},
            files={"file": ("lecture.pdf", pdf_bytes, "application/pdf")},
        )
        self.assertEqual(response.status_code, 201, response.text)
        document = response.json()
        self.assertEqual(document["status"], "ready")
        self.assertGreater(document["chunk_count"], 0)
        self.assertEqual(document["embedding_status"], "not_started")
        self.assertEqual(document["embedded_chunk_count"], 0)
        self.assertNotIn("stored_file_path", document)
        document_id = document["id"]

        list_response = self.client.get("/rag/documents")
        self.assertEqual(list_response.status_code, 200)
        self.assertEqual(len(list_response.json()), 1)

        patch_response = self.client.patch(
            f"/rag/documents/{document_id}",
            json={"title": "Yangilangan muhandislik grafikasi", "task_id": None},
        )
        self.assertEqual(patch_response.status_code, 200, patch_response.text)
        self.assertEqual(patch_response.json()["title"], "Yangilangan muhandislik grafikasi")
        self.assertIsNone(patch_response.json()["task_id"])

        empty_patch = self.client.patch(f"/rag/documents/{document_id}", json={})
        self.assertEqual(empty_patch.status_code, 422)

        chunks_response = self.client.get(f"/rag/documents/{document_id}/chunks")
        self.assertEqual(chunks_response.status_code, 200)
        chunks = chunks_response.json()
        self.assertEqual(len(chunks), document["chunk_count"])
        self.assertEqual(chunks[0]["page_number_start"], 1)
        self.assertNotIn("embedding", chunks[0])

        reprocess_response = self.client.post(f"/rag/documents/{document_id}/reprocess")
        self.assertEqual(reprocess_response.status_code, 200, reprocess_response.text)
        self.assertEqual(reprocess_response.json()["status"], "ready")

        delete_response = self.client.delete(f"/rag/documents/{document_id}")
        self.assertEqual(delete_response.status_code, 204)
        self.assertEqual(self.client.get("/rag/documents").json(), [])

    def test_duplicate_and_fake_pdf_are_rejected(self) -> None:
        pdf_bytes = make_pdf_bytes("Takroriy material")
        first = self.client.post(
            "/rag/documents",
            data={"title": "Birinchi material"},
            files={"file": ("same.pdf", pdf_bytes, "application/pdf")},
        )
        self.assertEqual(first.status_code, 201, first.text)

        duplicate = self.client.post(
            "/rag/documents",
            data={"title": "Takroriy material"},
            files={"file": ("same.pdf", pdf_bytes, "application/pdf")},
        )
        self.assertEqual(duplicate.status_code, 409)

        fake = self.client.post(
            "/rag/documents",
            data={"title": "Soxta PDF"},
            files={"file": ("fake.pdf", b"not a real pdf", "application/pdf")},
        )
        self.assertEqual(fake.status_code, 400)


if __name__ == "__main__":
    unittest.main()
